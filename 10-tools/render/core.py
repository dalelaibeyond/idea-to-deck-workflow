"""Core primitives for the idea-to-deck rendering engine.

Deterministic layout primitives shared across archetype renderers.
All geometry/theme data lives in ``registry.json``; the engine only
interprets data and never inlines deck content (D1 discipline).
"""

from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR


def hex_to_rgb(hex_str):
    hex_str = hex_str.lstrip("#")
    return RGBColor(*(int(hex_str[i:i + 2], 16) for i in (0, 2, 4)))


def resolve_color(color_key, theme):
    return theme[color_key] if color_key in theme else color_key


class Ctx:
    """Render context: registry data + derived helpers, shared per build."""

    def __init__(self, registry):
        self.registry = registry
        self.theme = registry["theme"]
        self.fonts = registry["fonts"]
        self.textframe = registry["common"]["text_frame"]
        self.bullet = registry["common"]["bullet"]
        self.common_geom = registry["common"]

    # -- primitives -------------------------------------------------------

    def fill_rect(self, slide, box, fill, radius=None):
        from pptx.enum.shapes import MSO_SHAPE
        shape = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE,
            Inches(box[0]), Inches(box[1]), Inches(box[2]), Inches(box[3]))
        if radius is not None:
            try:
                shape.adjustments[0] = radius
            except Exception:
                pass
        shape.fill.solid()
        shape.fill.fore_color.rgb = hex_to_rgb(resolve_color(fill, self.theme))
        shape.line.fill.background()
        shape.shadow.inherit = False
        return shape

    def add_text(self, slide, box, text, size, color, bold=False,
                 align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, line_spacing=None,
                 italic=False):
        tf_cfg = self.textframe
        box_obj = slide.shapes.add_textbox(
            Inches(box[0]), Inches(box[1]), Inches(box[2]), Inches(box[3]))
        tf = box_obj.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = anchor
        tf.margin_left = Inches(tf_cfg["margin_left"])
        tf.margin_right = Inches(tf_cfg["margin_right"])
        tf.margin_top = Inches(tf_cfg["margin_top"])
        tf.margin_bottom = Inches(tf_cfg["margin_bottom"])
        p = tf.paragraphs[0]
        p.alignment = align
        p.line_spacing = line_spacing if line_spacing is not None else tf_cfg["line_spacing_default"]
        run = p.add_run()
        run.text = text
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.italic = italic
        run.font.color.rgb = hex_to_rgb(resolve_color(color, self.theme))
        run.font.name = self.fonts["header"] if bold else self.fonts["body"]
        return box_obj

    def add_bullets(self, slide, box, items, size=None, color="text_primary",
                    bold_lead=None, line_spacing=None):
        tf_cfg = self.textframe
        bullet_cfg = self.bullet
        box_obj = slide.shapes.add_textbox(
            Inches(box[0]), Inches(box[1]), Inches(box[2]), Inches(box[3]))
        tf = box_obj.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(tf_cfg["margin_left"])
        tf.margin_right = Inches(tf_cfg["margin_right"])
        base_color = resolve_color(color, self.theme)
        size = size if size is not None else bullet_cfg["size_default"]
        bold_lead = bullet_cfg["bold_lead_default"] if bold_lead is None else bold_lead
        line_spacing = line_spacing if line_spacing is not None else tf_cfg["line_spacing_default"]
        prefix = bullet_cfg["prefix"]
        first = True
        for item in items:
            p = tf.paragraphs[0] if first else tf.add_paragraph()
            first = False
            p.line_spacing = line_spacing
            p.space_after = Pt(bullet_cfg["space_after"])
            lead, _, body = item.partition(":")
            if body:
                run = p.add_run()
                run.text = prefix
                run.font.size = Pt(size)
                run.font.color.rgb = hex_to_rgb(base_color)
                r = p.add_run()
                r.text = lead + ": "
                r.font.size = Pt(size)
                r.font.bold = bold_lead
                r.font.color.rgb = hex_to_rgb(base_color)
                r.font.name = self.fonts["body"]
                r2 = p.add_run()
                r2.text = body
                r2.font.size = Pt(size)
                r2.font.bold = False
                r2.font.color.rgb = hex_to_rgb(base_color)
                r2.font.name = self.fonts["body"]
            else:
                run = p.add_run()
                run.text = prefix + lead
                run.font.size = Pt(size)
                run.font.color.rgb = hex_to_rgb(base_color)
                run.font.name = self.fonts["body"]
        return box_obj

    def set_notes(self, slide, voiceover, mindset):
        notes = slide.notes_slide.notes_text_frame
        notes.text = "[Voiceover]\n{0}\n\n[Delivery Mindset]\n{1}".format(voiceover, mindset)

    # -- composite patterns -------------------------------------------------

    def slide_header(self, slide, title):
        geom = self.common_geom["slide_header"]
        self.fill_rect(slide, geom["top_rect"]["box"], geom["top_rect"]["fill"])
        tb = geom["title_box"]
        self.add_text(slide, tb["box"], title, tb["size"], tb["color"], bold=tb["bold"])

    def accent_strip(self, slide):
        geom = self.common_geom["accent_strip"]
        self.fill_rect(slide, geom["box"], geom["fill"])